"""Render a deck (from builder.build_deck) to a .pptx file with python-pptx."""

import tempfile
import uuid
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .media import (AUDIO_TYPES, VIDEO_TYPES, fetch_youtube_thumbnail, make_audio_icon,
                    make_poster)
from .themes import DEFAULT_THEME, THEMES

W, H = 13.333, 7.5          # 16:9 slide size in inches
MARGIN = 0.8
NO_FOOTER = ("title", "section", "closing")


def render_pptx(deck, out_path, theme_key=DEFAULT_THEME, media_dir=None):
    """media_dir is the folder holding the deck's uploaded videos, photos and audio."""
    theme = THEMES.get(theme_key, THEMES[DEFAULT_THEME])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]

    total = len(deck["slides"])
    with tempfile.TemporaryDirectory() as tmp:  # generated posters and icons
        for index, data in enumerate(deck["slides"], start=1):
            slide = prs.slides.add_slide(blank)
            r = _Slide(slide, theme, deck, media_dir=media_dir, tmp=Path(tmp))
            getattr(r, "layout_" + data["layout"])(data)
            full_photo = data["layout"] == "photo" and data.get("style") == "full"
            if data["layout"] not in NO_FOOTER and not full_photo:
                # on "photo left" slides keep the footer on the text side
                left = W * 0.48 + 0.7 if data.get("style") == "left" else MARGIN
                r.footer(deck["title"], index, total, left=left)
            # background music sits on the first slide, slide audio on its own slide
            audio = deck.get("music") if index == 1 else data.get("audio")
            if audio:
                r.audio(audio)
            if data.get("notes"):
                slide.notes_slide.notes_text_frame.text = data["notes"]
        prs.save(out_path)
    return out_path


def _rgb(hex_color):
    return RGBColor.from_string(hex_color)


class _Slide:
    def __init__(self, slide, theme, deck=None, media_dir=None, tmp=None):
        self.s, self.t, self.deck = slide, theme, deck or {}
        self.media_dir = Path(media_dir) if media_dir else None
        self.tmp = tmp

    def media(self, name):
        """Path of an uploaded media file, or None if it's missing."""
        if name and self.media_dir and (self.media_dir / name).is_file():
            return self.media_dir / name
        return None

    def picture_fill(self, path, x, y, w, h):
        """Add a picture that fills the box, cropping (not stretching) the image."""
        from PIL import Image

        with Image.open(path) as im:
            iw, ih = im.size
        pic = self.s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        box, img = w / h, iw / ih
        if img > box:
            pic.crop_left = pic.crop_right = (1 - box / img) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - img / box) / 2
        return pic

    def audio(self, name):
        """Embed an audio file as a speaker icon in the bottom-right corner."""
        path = self.media(name)
        if not path:
            return
        icon = make_audio_icon(self.t, self.tmp / f"{uuid.uuid4().hex}.png")
        size = 0.6
        shp = self.s.shapes.add_movie(
            str(path), Inches(W - MARGIN - size), Inches(H - 1.25), Inches(size), Inches(size),
            poster_frame_image=str(icon),
            mime_type=AUDIO_TYPES.get(path.suffix.lower(), "audio/mpeg"))
        # python-pptx only knows video; re-tag it so PowerPoint treats it as sound
        media = shp._element.find(".//" + qn("a:videoFile"))
        if media is not None:
            media.tag = qn("a:audioFile")
            rel = self.s.part.rels[media.get(qn("r:link"))]
            rel._reltype = rel.__dict__["reltype"] = (
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio")

    # ------------------------------------------------------------ primitives

    def background(self, color):
        fill = self.s.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color)

    def gradient_background(self):
        fill = self.s.background.fill
        fill.gradient()
        fill.gradient_angle = 135
        stops = fill.gradient_stops
        # PowerPoint measures the angle the other way round from CSS, so the
        # stops are swapped to get hero1 in the top-left like the web preview.
        stops[0].color.rgb = _rgb(self.t["hero2"])
        stops[1].color.rgb = _rgb(self.t["hero1"])

    def shape(self, kind, x, y, w, h, color, alpha=None, line=False):
        shp = self.s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(color)
        if alpha is not None:  # python-pptx has no API for transparency, set it in XML
            clr = shp.fill._xPr.find(qn("a:solidFill"))[0]
            el = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
            clr.append(el)
        if not line:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def text(self, x, y, w, h, text, size=18, color=None, bold=False, font=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
        box = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        self._style(run, size, color or self.t["text"], bold, font or self.t["body_font"], italic)
        return box

    @staticmethod
    def _style(run, size, color, bold, font, italic=False):
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = _rgb(color)

    def heading(self, title):
        """Standard slide heading: accent bar + title."""
        self.shape(MSO_SHAPE.RECTANGLE, MARGIN, 0.72, 0.09, 0.62, self.t["accent"])
        size = 32 if len(title) < 45 else 26
        self.text(MARGIN + 0.25, 0.6, W - 2 * MARGIN - 0.25, 0.9, title, size=size,
                  bold=True, font=self.t["title_font"], anchor=MSO_ANCHOR.MIDDLE)

    def footer(self, deck_title, index, total, left=MARGIN):
        self.text(left, H - 0.55, W - left - MARGIN - 2.2, 0.35, deck_title, size=10,
                  color=self.t["muted"])
        self.text(W - MARGIN - 2, H - 0.55, 2, 0.35, f"{index} / {total}", size=10,
                  color=self.t["muted"], align=PP_ALIGN.RIGHT)

    def decorations(self):
        """Soft translucent circles used on hero slides."""
        white = "FFFFFF"
        self.shape(MSO_SHAPE.OVAL, W - 4.6, -1.8, 6.5, 6.5, white, alpha=0.07)
        self.shape(MSO_SHAPE.OVAL, W - 2.6, 3.9, 4.2, 4.2, white, alpha=0.06)
        self.shape(MSO_SHAPE.OVAL, -1.4, H - 1.6, 3.0, 3.0, white, alpha=0.05)

    # --------------------------------------------------------------- layouts

    def layout_title(self, d):
        t = self.t
        cover = self.media(self.deck.get("cover"))
        if cover:  # photo cover: the image with a tinted overlay so text stays readable
            self.picture_fill(cover, 0, 0, W, H)
            self.shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H, t["hero1"], alpha=0.62)
        else:
            self.gradient_background()
            self.decorations()
        self.shape(MSO_SHAPE.RECTANGLE, MARGIN, 2.35, 1.2, 0.08, t["hero_text"])
        size = 54 if len(d["title"]) < 40 else 42 if len(d["title"]) < 70 else 34
        self.text(MARGIN, 2.6, W - 4, 2.4, d["title"], size=size, bold=True,
                  color=t["hero_text"], font=t["title_font"], anchor=MSO_ANCHOR.TOP)
        if d.get("subtitle"):
            self.text(MARGIN, 5.0, W - 4.5, 1.2, d["subtitle"], size=20,
                      color=t["hero_text"])
        self.text(MARGIN, H - 0.9, 6, 0.4, date.today().strftime("%B %Y"), size=12,
                  color=t["hero_text"])

    def layout_agenda(self, d):
        t = self.t
        self.background(t["bg"])
        panel = self.shape(MSO_SHAPE.RECTANGLE, 0, 0, 4.2, H, t["hero1"])
        panel.fill.gradient()
        panel.fill.gradient_angle = 90
        panel.fill.gradient_stops[0].color.rgb = _rgb(t["hero2"])
        panel.fill.gradient_stops[1].color.rgb = _rgb(t["hero1"])
        self.text(0.7, 2.9, 3.2, 1.2, d["title"], size=44, bold=True,
                  color=t["hero_text"], font=t["title_font"])
        items = d["items"]
        cols = 2 if len(items) > 5 else 1
        per_col = -(-len(items) // cols)
        col_w = (W - 4.2 - 1.4) / cols
        row_h = min(1.05, (H - 2.2) / per_col)
        top = (H - row_h * per_col) / 2
        for i, item in enumerate(items):
            c, r = divmod(i, per_col)
            x, y = 5.0 + c * col_w, top + r * row_h
            self.text(x, y, 0.9, row_h, f"{i + 1:02d}", size=26, bold=True,
                      color=t["accent"], font=t["title_font"], anchor=MSO_ANCHOR.MIDDLE)
            self.text(x + 0.95, y, col_w - 1.1, row_h, item, size=22 if cols == 1 else 17,
                      anchor=MSO_ANCHOR.MIDDLE)

    def layout_section(self, d):
        t = self.t
        self.gradient_background()
        self.decorations()
        self.text(MARGIN, 1.2, 5, 2.2, d["number"], size=120, bold=True,
                  color=t["hero_text"], font=t["title_font"])
        self.shape(MSO_SHAPE.RECTANGLE, MARGIN + 0.1, 3.75, 1.0, 0.07, t["hero_text"])
        self.text(MARGIN, 3.95, W - 3, 1.5, d["title"], size=44, bold=True,
                  color=t["hero_text"], font=t["title_font"])
        if d.get("subtitle"):
            self.text(MARGIN, 5.35, W - 4, 1.2, d["subtitle"], size=18, color=t["hero_text"])

    def layout_bullets(self, d):
        t = self.t
        self.background(t["bg"])
        self.heading(d["title"])
        bullets = d["bullets"]
        chars = sum(len(b) for b in bullets)
        size = 22 if chars < 260 else 19 if chars < 420 else 16
        top, avail = 1.85, H - 1.85 - 0.9
        row_h = avail / max(len(bullets), 3)
        for i, b in enumerate(bullets):
            y = top + i * row_h
            self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN + 0.25, y + 0.17, 0.2, 0.2,
                       t["accent"] if i % 2 == 0 else t["accent2"])
            self.text(MARGIN + 0.7, y, W - 2 * MARGIN - 0.9, row_h, b, size=size)
        # thin decorative strip on the right edge
        self.shape(MSO_SHAPE.RECTANGLE, W - 0.18, 0, 0.18, H, t["accent"])

    def layout_cards(self, d):
        t = self.t
        self.background(t["bg"])
        self.heading(d["title"])
        cards = d["cards"]
        cols = len(cards) if len(cards) <= 3 else (2 if len(cards) == 4 else 3)
        rows = -(-len(cards) // cols)
        gap = 0.35
        cw = (W - 2 * MARGIN - gap * (cols - 1)) / cols
        top = 1.9
        ch = min(3.8, (H - top - 0.85 - gap * (rows - 1)) / rows)
        top += (H - 0.85 - top - (ch * rows + gap * (rows - 1))) / 2  # centre vertically
        one_row = rows == 1
        for i, card in enumerate(cards):
            r, c = divmod(i, cols)
            x, y = MARGIN + c * (cw + gap), top + r * (ch + gap)
            box = self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, ch, t["bg2"])
            box.adjustments[0] = 0.06
            color = t["accent"] if i % 2 == 0 else t["accent2"]
            self.shape(MSO_SHAPE.RECTANGLE, x + 0.35, y + 0.4, 0.7, 0.08, color)
            # heading and body share one text box so a wrapped heading pushes the body down
            box = self.text(x + 0.35, y + 0.6, cw - 0.7, ch - 0.8, card["head"],
                            size=24 if one_row else 18, bold=True, color=color,
                            font=t["title_font"])
            body = box.text_frame.add_paragraph()
            body.space_before = Pt(10)
            self._style(body.add_run(), 18 if one_row else 14, t["text"], False, t["body_font"])
            body.runs[0].text = card["body"]

    def layout_stats(self, d):
        t = self.t
        self.background(t["bg"])
        self.heading(d["title"])
        stats = d["stats"]
        gap = 0.4
        cw = (W - 2 * MARGIN - gap * (len(stats) - 1)) / len(stats)
        for i, st in enumerate(stats):
            x = MARGIN + i * (cw + gap)
            box = self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.2, cw, 3.9, t["bg2"])
            box.adjustments[0] = 0.06
            color = t["accent"] if i % 2 == 0 else t["accent2"]
            self.text(x, 2.6, cw, 1.6, st["value"], size=60 if len(st["value"]) < 6 else 44,
                      bold=True, color=color, font=t["title_font"], align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE)
            self.shape(MSO_SHAPE.RECTANGLE, x + cw / 2 - 0.4, 4.3, 0.8, 0.06, color)
            self.text(x + 0.3, 4.5, cw - 0.6, 1.4, st["label"], size=17,
                      color=t["text"], align=PP_ALIGN.CENTER)

    def layout_table(self, d):
        t = self.t
        self.background(t["bg"])
        self.heading(d["title"])
        rows = d["rows"]
        n_rows, n_cols = len(rows), len(rows[0])
        row_h = min(0.75, (H - 2.9) / n_rows)
        shape = self.s.shapes.add_table(n_rows, n_cols, Inches(MARGIN), Inches(1.9),
                                        Inches(W - 2 * MARGIN), Inches(row_h * n_rows))
        table = shape.table
        # remove the default PowerPoint table style banding
        tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
        tbl_pr.set("bandRow", "0")
        tbl_pr.set("firstRow", "0")
        size = 18 if n_cols <= 4 else 14
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = _rgb(t["accent"])
                else:
                    cell.fill.fore_color.rgb = _rgb(t["bg2"] if r % 2 else t["bg"])
                cell.margin_left = cell.margin_right = Inches(0.15)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = value
                self._style(run, size, "FFFFFF" if r == 0 else t["text"], r == 0,
                            t["body_font"])

    def layout_statement(self, d):
        t = self.t
        self.background(t["bg2"])
        self.shape(MSO_SHAPE.RECTANGLE, 0, 0, 0.25, H, t["accent"])
        self.text(MARGIN + 0.4, 0.9, 2, 1.6, "“", size=140, bold=True,
                  color=t["accent"], font="Georgia")
        size = 34 if len(d["text"]) < 120 else 28
        self.text(MARGIN + 0.6, 2.3, W - 2 * MARGIN - 1.2, 3.4, d["text"], size=size,
                  font=t["title_font"], anchor=MSO_ANCHOR.MIDDLE)
        self.text(MARGIN + 0.6, 5.9, W - 3, 0.5, "— " + d["title"], size=16,
                  color=t["muted"], italic=True)

    def layout_video(self, d):
        """Uploaded files are embedded and play inside PowerPoint;
        links become a clickable cover image that opens the video."""
        t = self.t
        self.background(t["bg"])
        self.heading(d["title"])
        vh = 4.5 if d.get("caption") else 4.75
        vw = vh * 16 / 9
        x, y = (W - vw) / 2, 1.75
        self.shape(MSO_SHAPE.RECTANGLE, x + 0.12, y + 0.12, vw, vh, t["accent"], alpha=0.35)

        poster = self.tmp / f"{uuid.uuid4().hex}.png"
        video = self.media_dir / d["file"] if d.get("file") and self.media_dir else None
        if video and video.is_file():
            make_poster(t, poster)
            self.s.shapes.add_movie(str(video), Inches(x), Inches(y), Inches(vw), Inches(vh),
                                    poster_frame_image=str(poster),
                                    mime_type=VIDEO_TYPES.get(video.suffix.lower(), "video/mp4"))
        else:
            url = d.get("url", "")
            thumb = poster.with_suffix(".jpg")
            if d.get("youtube") and fetch_youtube_thumbnail(d["youtube"], thumb):
                pic = self.s.shapes.add_picture(str(thumb), Inches(x), Inches(y), Inches(vw), Inches(vh))
                self._play_button(x + vw / 2, y + vh / 2, url)
            else:
                pic = self.s.shapes.add_picture(str(make_poster(t, poster)), Inches(x), Inches(y),
                                                Inches(vw), Inches(vh))
            if url:
                pic.click_action.hyperlink.address = url
                # link on the shape, not the run, so PowerPoint keeps the theme colour
                link = self.text(x, y + vh + 0.1, vw, 0.4,
                                 "▶  Click the video to watch" + (" on YouTube" if d.get("youtube") else ""),
                                 size=13, color=t["muted"], align=PP_ALIGN.CENTER)
                link.click_action.hyperlink.address = url

        if d.get("caption"):
            self.text(x, y + vh + (0.5 if d.get("url") else 0.15), vw, 0.5, d["caption"],
                      size=16, color=t["muted"], align=PP_ALIGN.CENTER)

    def layout_photo(self, d):
        """style: full (edge to edge), left / right (photo beside text), framed (photo + caption)."""
        t = self.t
        style = d.get("style", "full")
        path = self.media(d.get("image"))
        lines = [l.strip() for l in d.get("text", "").splitlines() if l.strip()]

        if style == "full":
            self.background(t["hero1"])
            if path:
                self.picture_fill(path, 0, 0, W, H)
            self.shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H, "000000", alpha=0.25)
            self.shape(MSO_SHAPE.RECTANGLE, 0, H * 0.52, W, H * 0.48, "000000", alpha=0.45)
            self.shape(MSO_SHAPE.RECTANGLE, MARGIN, 4.55, 1.0, 0.07, t["accent"])
            self.text(MARGIN, 4.7, W - 2 * MARGIN, 1.0, d["title"], size=40, bold=True,
                      color="FFFFFF", font=t["title_font"])
            if lines:
                self.text(MARGIN, 5.75, W - 3, 1.2, " ".join(lines), size=20, color="FFFFFF")
            return

        if style == "framed":
            self.background(t["bg"])
            self.heading(d["title"])
            vh = 4.4 if lines else 4.85
            vw = min(vh * 16 / 9, W - 2 * MARGIN)
            x, y = (W - vw) / 2, 1.7
            self.shape(MSO_SHAPE.RECTANGLE, x + 0.12, y + 0.12, vw, vh, t["accent"], alpha=0.35)
            if path:
                self.picture_fill(path, x, y, vw, vh)
            if lines:
                self.text(x, y + vh + 0.2, vw, 0.5, " ".join(lines), size=16,
                          color=t["muted"], align=PP_ALIGN.CENTER)
            return

        # left / right: photo fills one half, text on the other
        self.background(t["bg"])
        half = W * 0.48
        px = 0 if style == "left" else W - half
        if path:
            self.picture_fill(path, px, 0, half, H)
        tx = half + 0.7 if style == "left" else MARGIN
        tw = W - half - 0.7 - MARGIN
        self.shape(MSO_SHAPE.RECTANGLE, tx, 1.3, 0.9, 0.08, t["accent"])
        self.text(tx, 1.5, tw, 1.6, d["title"], size=34 if len(d["title"]) < 40 else 28,
                  bold=True, font=t["title_font"])
        if len(lines) > 1:  # several lines become bullets
            row = min(0.85, 3.4 / len(lines))
            for i, line in enumerate(lines[:6]):
                y = 3.2 + i * row
                self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, y + 0.13, 0.16, 0.16,
                           t["accent"] if i % 2 == 0 else t["accent2"])
                self.text(tx + 0.35, y, tw - 0.35, row, line, size=18)
        elif lines:
            self.text(tx, 3.2, tw, 3.2, lines[0], size=20, color=t["muted"])

    def _play_button(self, cx, cy, url):
        r = 0.55
        circle = self.shape(MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, "FFFFFF", alpha=0.92)
        tri = self.shape(MSO_SHAPE.ISOSCELES_TRIANGLE, cx - 0.22, cy - 0.25, 0.5, 0.5, self.t["hero1"])
        tri.rotation = 90
        for shp in (circle, tri):
            shp.click_action.hyperlink.address = url

    def layout_closing(self, d):
        t = self.t
        self.gradient_background()
        self.decorations()
        self.text(0, 2.5, W, 1.5, d["title"], size=66, bold=True, color=t["hero_text"],
                  font=t["title_font"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        self.shape(MSO_SHAPE.RECTANGLE, W / 2 - 0.6, 4.15, 1.2, 0.07, t["hero_text"])
        self.text(0, 4.4, W, 0.8, d["subtitle"], size=22, color=t["hero_text"],
                  align=PP_ALIGN.CENTER)
