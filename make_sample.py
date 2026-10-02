"""Create samples/project_report.docx - a demo document to try SlideGen with."""

from pathlib import Path

import docx

doc = docx.Document()
doc.add_heading("Smart Farming Initiative 2026", level=0)
doc.add_paragraph(
    "A plan to help 5,000 smallholder farmers raise yields with affordable sensors and mobile advice. "
    "This report summarises the problem, our solution, the results of the pilot and the next steps."
)

doc.add_heading("The Challenge", level=1)
doc.add_paragraph(
    "Most smallholder farmers still rely on guesswork for irrigation and fertiliser. "
    "Weather is becoming less predictable every season. "
    "Extension officers can visit each farm only once or twice a year. "
    "As a result, harvests vary widely and incomes remain low."
)

doc.add_heading("Our Solution", level=1)
doc.add_paragraph("We combine three simple components into one service that any farmer can use from a basic phone.")
doc.add_heading("Key Components", level=2)
for item in [
    "Soil sensors: low-cost probes measure moisture and temperature every hour.",
    "SMS advisor: farmers receive daily, plain-language advice in Kinyarwanda, English or French.",
    "Cooperative dashboard: cooperatives see the status of every member farm in one place.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("How It Works", level=2)
for item in [
    "Sensors send readings to the cloud over the mobile network",
    "The platform combines readings with local weather forecasts",
    "A recommendation engine decides when to irrigate and fertilise",
    "Farmers receive a short SMS with the action to take",
    "Results are tracked so advice improves every season",
]:
    doc.add_paragraph(item, style="List Number")

doc.add_heading("Pilot Results", level=1)
doc.add_paragraph("The six-month pilot ran with 320 farmers in three districts.")
doc.add_heading("Impact", level=2)
for item in [
    "32% increase in average maize yield",
    "40% less water used for irrigation",
    "320 farmers enrolled in the pilot",
    "4.6 average satisfaction score out of 5",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Results by District", level=2)
rows = [
    ("District", "Farmers", "Yield Gain", "Water Saved"),
    ("Bugesera", "120", "+35%", "44%"),
    ("Nyagatare", "110", "+30%", "38%"),
    ("Huye", "90", "+29%", "37%"),
]
table = doc.add_table(rows=len(rows), cols=4)
for r, row in enumerate(rows):
    for c, value in enumerate(row):
        table.cell(r, c).text = value

doc.add_heading("Vision", level=1)
doc.add_paragraph("Every farmer deserves the same quality of information as the largest commercial farms.")

doc.add_heading("Next Steps", level=1)
for item in [
    "Expand to 5,000 farmers across eight districts by June 2027",
    "Partner with two mobile network operators to reduce SMS costs",
    "Add pest and disease alerts using photos sent by farmers",
    "Train 40 cooperative leaders as local champions",
]:
    doc.add_paragraph(item, style="List Bullet")

out = Path(__file__).parent / "samples"
out.mkdir(exist_ok=True)
doc.save(out / "project_report.docx")
print("Saved", out / "project_report.docx")
